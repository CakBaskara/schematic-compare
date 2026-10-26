import collections
import io
import os
import re
import zipfile

import fitz
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage

HDR = "Times New Roman"
DAT = "Arial"
_thin = Side(style="thin")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


class _Renderer:
    def __init__(self):
        self._docs = {}

    def png(self, path, page, rect, box_w, box_h, oversample=1.3, colors=96):
        doc = self._docs.get(path) or self._docs.setdefault(path, fitz.open(path))
        x0, y0, x1, y1 = rect
        s = min(box_w / max(x1 - x0, 1), box_h / max(y1 - y0, 1))
        dpi = max(72, int(round(s * 72 * oversample)))
        pix = doc[page - 1].get_pixmap(dpi=dpi, clip=fitz.Rect(*rect))
        img = PILImage.frombytes("RGB", (pix.width, pix.height), pix.samples)
        img = img.quantize(colors=colors, method=PILImage.Quantize.MEDIANCUT)
        buf = io.BytesIO()
        img.save(buf, format="PNG", optimize=True)
        buf.seek(0)
        return buf, int((x1 - x0) * s), int((y1 - y0) * s)


def _place(ws, buf, w, h, col, row, off=6):
    im = XLImage(buf)
    im.width, im.height = w, h
    marker = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(off), row=row - 1, rowOff=pixels_to_EMU(off))
    im.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
    ws.add_image(im)


def _col_px(width):
    return int(width * 7 + 5)


def _style(ws, rng, font=None, align=None, fill=None, border=BORDER):
    for row in ws[rng]:
        for c in row:
            if font:
                c.font = font
            if align:
                c.alignment = align
            if fill:
                c.fill = fill
            if border:
                c.border = border


def _table(ws, r, cols, rows, widths_merge, header_font=None):
    """Write a small boxed table; widths_merge gives the column span of every logical column."""
    starts, c = [], 2
    for span in widths_merge:
        starts.append(c)
        c += span
    last = c - 1
    for i, values in enumerate([cols] + rows):
        for v, s, span in zip(values, starts, widths_merge):
            ws.cell(r + i, s, v)
            if span > 1:
                ws.merge_cells(start_row=r + i, start_column=s, end_row=r + i, end_column=s + span - 1)
        head = i == 0
        _style(ws, f"B{r + i}:{get_column_letter(last)}{r + i}",
               font=header_font if head and header_font else Font(name=HDR if head else DAT, size=10, bold=head),
               align=Alignment(horizontal="center" if head else "left", vertical="center", wrap_text=True),
               fill=HEAD_FILL if head else None)
    return r + len(rows) + 1


def write(path, ctx):
    cached = {}
    rnd = _Renderer()
    doc_b, doc_a = ctx["docs"]
    name_b, name_a = ctx["names"]
    rows = ctx["rows"]

    wb = openpyxl.Workbook()
    ov = wb.active
    ov.title = "Overview"
    cpl = wb.create_sheet("Change Point List")
    cdc = wb.create_sheet("Circuit Diagram Change")

    # Overview ---------------------------------------------------------------
    ov.sheet_view.showGridLines = False
    ov.sheet_view.zoomScale = 85
    ov.column_dimensions["A"].width = 2
    for c in range(2, 20):
        ov.column_dimensions[get_column_letter(c)].width = 13
    ov["B2"] = f"{name_b} ({doc_b.stem})"
    ov["K2"] = f"{name_a} ({doc_a.stem})"
    ov.merge_cells("B2:I2")
    ov.merge_cells("K2:S2")
    _style(ov, "B2:I2", font=Font(name=HDR, size=12, bold=True), align=CENTER)
    _style(ov, "K2:S2", font=Font(name=HDR, size=12, bold=True), align=CENTER)
    shown = []
    for doc, col, span in ((doc_b, 2, 8), (doc_a, 11, 9)):
        page = next((p for p in doc.pages if "block" in p.sheet_title.lower()), doc.pages[0])
        buf, w, h = rnd.png(doc.path, page.number, page.border, _col_px(13) * span - 12, 600)
        _place(ov, buf, w, h, col, 3)
        shown.append((h, page))
    r = 3 + int(max(h for h, _ in shown) / 20) + 2
    ov.cell(r, 2, "Picture: " + " / ".join(p.sheet_title or f"sheet {p.number:02d}" for _, p in shown)).font = Font(name=DAT, size=9, italic=True, color="808080")
    r += 2

    def created(doc):
        m = re.match(r"D:(\d{4})(\d{2})(\d{2})", doc.meta.get("creationDate") or "")
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else "-"

    p1b, p1a = doc_b.pages[0], doc_a.pages[0]
    ov.cell(r, 2, "A. DOCUMENTS").font = Font(name=HDR, size=11, bold=True)
    r = _table(ov, r + 1, ("Item", f"{name_b} (Before)", f"{name_a} (After)"), [
        ("PDF file", doc_b.name, doc_a.name),
        ("Title block", f"{p1b.title or '-'}  Rev {p1b.rev or '-'}", f"{p1a.title or '-'}  Rev {p1a.rev or '-'}"),
        ("Sheet date (sheet 01)", p1b.date or "-", p1a.date or "-"),
        ("PDF created", created(doc_b), created(doc_a)),
        ("Sheets", str(len(doc_b.pages)), str(len(doc_a.pages))),
    ], (2, 7, 9)) + 1

    ov.cell(r, 2, "B. SHEET MAPPING").font = Font(name=HDR, size=11, bold=True)
    n = max(len(doc_b.pages), len(doc_a.pages))
    mapping = []
    for i in range(1, n + 1):
        tb = doc_b.pages[i - 1].sheet_title if i <= len(doc_b.pages) else "-"
        ta = doc_a.pages[i - 1].sheet_title if i <= len(doc_a.pages) else "-"
        mapping.append((f"{i:02d}", tb or "-", ta or "-"))
    r = _table(ov, r + 1, ("Sheet", f"{name_b} (Before)", f"{name_a} (After)"), mapping, (2, 7, 9)) + 1

    ov.cell(r, 2, "C. SUMMARY (from sheet 'Change Point List')").font = Font(name=HDR, size=11, bold=True)
    last = 10 + len(rows)
    rem = f"'Change Point List'!$K$11:$K${last}"
    mnt = f"'Change Point List'!$J$11:$J${last}"
    kinds = collections.Counter(x.kind for x in rows)
    mounts = collections.Counter(x.mount for x in rows)
    summary = [
        ("Listed parts", f"=COUNTA('Change Point List'!$E$11:$E${last})", len(rows)),
        ("Added", f'=COUNTIF({rem},"Added*")', kinds["Added"]),
        ("Removed", f'=COUNTIF({rem},"Removed*")', kinds["Removed"]),
        ("Changed (value / footprint / NC)", f'=COUNTIF({rem},"Changed*")', kinds["Changed"]),
        ("Moved only (same part, other circuit)", f'=COUNTIF({rem},"Moved:*")', kinds["Moved"]),
        ("Mounting After = Mounted", f'=COUNTIF({mnt},"Mounted")', mounts["Mounted"]),
        ("Mounting After = Not Mounted", f'=COUNTIF({mnt},"Not Mounted")', mounts["Not Mounted"]),
    ]
    r += 1
    ov.cell(r, 2, "Item")
    ov.cell(r, 6, "Count")
    ov.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    _style(ov, f"B{r}:F{r}", font=Font(name=HDR, size=10, bold=True), align=CENTER, fill=HEAD_FILL)
    for i, (label, formula, value) in enumerate(summary, start=1):
        ov.cell(r + i, 2, label)
        ov.merge_cells(start_row=r + i, start_column=2, end_row=r + i, end_column=5)
        ov.cell(r + i, 6, formula)
        cached[("Overview", f"F{r + i}")] = value
        _style(ov, f"B{r + i}:F{r + i}", font=Font(name=DAT, size=10), align=Alignment(horizontal="left"))
        ov.cell(r + i, 6).alignment = Alignment(horizontal="center")
    r += len(summary) + 2

    ov.cell(r, 2, "D. NOTES").font = Font(name=HDR, size=11, bold=True)
    notes = [
        "Generated by sch_compare from the two PDFs. Parts are compared by reference designator (Ref); value / part name and footprint are as printed. Parts with the same Ref, value, footprint and circuit are not listed.",
        f"Prefixes {', '.join(ctx['exclude'])} are not listed (test points). '-' in a footprint column means no footprint is printed for that part.",
        "'Moved' = same Ref and value now used in another circuit. 'Untitled' blocks have no heading on the sheet; they are named after their main part.",
        "Values are read from text positions. Check sheet 'Extraction Review' and fix misreads with a config file (see README).",
    ]
    for i, t in enumerate(notes, start=1):
        c = ov.cell(r + i, 2, f"{i}. {t}")
        c.font = Font(name=DAT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ov.merge_cells(start_row=r + i, start_column=2, end_row=r + i, end_column=19)
        ov.row_dimensions[r + i].height = 28
    r += len(notes) + 2

    if ctx.get("checks"):
        ov.cell(r, 2, "E. CHECK ITEMS").font = Font(name=HDR, size=11, bold=True)
        r = _table(ov, r + 1, ("No.", "Sheet", "Item"),
                   [(i, sh, t) for i, (sh, t) in enumerate(ctx["checks"], start=1)], (1, 2, 15))

    # Change Point List --------------------------------------------------------
    ws = cpl
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    for k, v in {"A": 3, "B": 6, "C": 7, "D": 30, "E": 9, "F": 34, "G": 17, "H": 36, "I": 17, "J": 13, "K": 62}.items():
        ws.column_dimensions[k].width = v
    ws["B4"] = f"Change Point List {name_b} and {name_a}"
    ws["B4"].font = Font(name=HDR, size=14)
    db, da = doc_b.date().strftime("%Y-%m-%d"), doc_a.date().strftime("%Y-%m-%d")
    ws["K5"] = f"Schematic: {name_b} ({db}) to {name_a} ({da})"
    ws["K5"].font = Font(name=HDR, size=12)
    ws["K5"].alignment = Alignment(horizontal="right")
    ws["B7"] = "A. CHANGE PART LIST"
    ws["B7"].font = Font(name=HDR, size=11, bold=True)
    for col, text in (("B", "No."), ("C", "Sheet"), ("D", "Circuit"), ("E", "Ref"), ("J", "Mounting\nAfter"), ("K", "Remarks")):
        ws[f"{col}8"] = text
        ws.merge_cells(f"{col}8:{col}10")
    ws["F8"] = f"{name_b} (Before)"
    ws.merge_cells("F8:G8")
    ws["H8"] = f"{name_a} (After)"
    ws.merge_cells("H8:I8")
    for col, text in (("F", "Part Name / Value"), ("G", "Footprint"), ("H", "Part Name / Value"), ("I", "Footprint")):
        ws[f"{col}9"] = text
        ws.merge_cells(f"{col}9:{col}10")
    _style(ws, "B8:K10", font=Font(name=HDR, size=11, bold=True), align=CENTER, fill=HEAD_FILL)
    ws.freeze_panes = "B11"
    dat = Font(name=DAT, size=10)
    for i, x in enumerate(rows):
        rr = 11 + i
        for c, v in enumerate((i + 1, x.sheet, x.circuit, x.ref, x.b_val, x.b_fp, x.a_val, x.a_fp), start=2):
            ws.cell(rr, c, v)
        ws.cell(rr, 10, f'=IF(OR(H{rr}="-",H{rr}="NC",LEFT(H{rr},3)="NC/"),"Not Mounted","Mounted")')
        cached[("Change Point List", f"J{rr}")] = x.mount
        ws.cell(rr, 11, x.remark)
        for c in range(2, 12):
            cell = ws.cell(rr, c)
            cell.font = dat
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if c == 11 else "center", vertical="center", wrap_text=c in (4, 6, 8, 11))

    def merge_runs(col, key):
        i = 0
        while i < len(rows):
            j = i
            while j + 1 < len(rows) and key(rows[j + 1]) == key(rows[i]):
                j += 1
            if j > i:
                ws.merge_cells(start_row=11 + i, start_column=col, end_row=11 + j, end_column=col)
            i = j + 1
    merge_runs(3, lambda x: x.sheet)
    merge_runs(4, lambda x: (x.sheet, x.circuit))
    ws.print_title_rows = "8:10"

    # Circuit Diagram Change ---------------------------------------------------
    ws = cdc
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 70
    img_w = 105
    for k, v in {"A": 4, "B": 8, "C": 18, "D": img_w, "E": img_w, "F": 48}.items():
        ws.column_dimensions[k].width = v
    for col, text in (("B", "SHEET"), ("C", "CIRCUIT"), ("D", f"BEFORE ({name_b.upper()})"), ("E", f"AFTER ({name_a.upper()})"), ("F", "REMARKS")):
        ws[f"{col}3"] = text
        ws.merge_cells(f"{col}3:{col}4")
    _style(ws, "B3:F4", font=Font(name=HDR, size=12, bold=True), align=CENTER, fill=HEAD_FILL)
    box_w, box_h = _col_px(img_w) - 14, 520
    rr = 5
    for b, a, remark in ctx["cdc"]:
        hs = []
        for blk, col, doc in ((b, 4, doc_b), (a, 5, doc_a)):
            if blk is None:
                c = ws.cell(rr, col, "Not Available")
                c.font = Font(name=DAT, size=14, color="808080")
                c.alignment = Alignment(horizontal="center", vertical="center")
                hs.append(60)
                continue
            buf, w, h = rnd.png(doc.path, blk.page, blk.crop or blk.region, box_w, box_h)
            _place(ws, buf, w, h, col, rr)
            hs.append(h + 12)
        ws.row_dimensions[rr].height = min(409, max(hs) * 0.75)
        if a and b and a.page != b.page:
            sheet = f"{b.page} to {a.page}"
        else:
            sheet = str((a or b).page)
        if a and b and a.name != b.name:
            circuit = f"{b.name} -> {a.name}"
        else:
            circuit = (a or b).name
        for c, v in ((2, sheet), (3, circuit), (6, remark)):
            ws.cell(rr, c, v)
        for c in range(2, 7):
            cell = ws.cell(rr, c)
            cell.border = BORDER
            if c in (2, 3):
                cell.font = Font(name=DAT, size=11)
                cell.alignment = CENTER
            elif c == 6:
                cell.font = Font(name=DAT, size=10)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        rr += 1
    ws.print_title_rows = "3:4"

    # Pin change sheets --------------------------------------------------------
    for pc in ctx.get("pins", []):
        ws = wb.create_sheet(f"Pin Change {pc['ref']}"[:31])
        ws.sheet_view.showGridLines = False
        for k, v in {"A": 3, "B": 7, "C": 28, "D": 28}.items():
            ws.column_dimensions[k].width = v
        ws["B2"] = f"{pc['ref']} pin changes ({name_b} sheet {pc['page_b']:02d} / {name_a} sheet {pc['page_a']:02d})"
        ws["B2"].font = Font(name=HDR, size=14)
        ws["B3"] = "Only pins whose net name changed are listed."
        ws["B3"].font = Font(name=DAT, size=9, italic=True, color="808080")
        _table(ws, 5, ("Pin", f"{name_b} (Before)", f"{name_a} (After)"), pc["rows"], (1, 1, 1),
               header_font=Font(name=HDR, size=11, bold=True))
        for row in ws.iter_rows(min_row=6, min_col=2, max_col=4):
            for c in row:
                c.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "B6"

    # Extraction review ---------------------------------------------------------
    ws = wb.create_sheet("Extraction Review")
    ws.sheet_view.showGridLines = False
    for k, v in {"A": 3, "B": 10, "C": 7, "D": 44, "E": 46, "F": 12}.items():
        ws.column_dimensions[k].width = v
    ws["B2"] = "Items to check by eye (text that could not be linked safely, values taken from far text)"
    ws["B2"].font = Font(name=HDR, size=12, bold=True)
    issue_rows = [(side, it["page"], it["kind"], it["text"], f"{it['x']}, {it['y']}")
                  for side, items in (("Before", ctx["issues"][0]), ("After", ctx["issues"][1])) for it in items]
    r = _table(ws, 4, ("Side", "Sheet", "Kind", "Text", "Position (pt)"), issue_rows or [("-", "-", "No issue", "-", "-")], (1, 1, 1, 1, 1))
    r += 1
    ws.cell(r, 2, "Blocks found (use 'block_names' in the config to rename, key = sheet:auto name)").font = Font(name=HDR, size=12, bold=True)
    brow = []
    for side, blocks in (("Before", ctx["blocks"][0]), ("After", ctx["blocks"][1])):
        for blk in sorted(blocks.values(), key=lambda x: (x.page, x.region[1], x.region[0])):
            brow.append((side, blk.page, blk.auto_name, blk.name, len(blk.refs)))
    _table(ws, r + 2, ("Side", "Sheet", "Auto name", "Name used", "Parts"), brow, (1, 1, 1, 1, 1))

    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    _inject_cached(path, cached)


def _inject_cached(path, cached):
    """Store formula results next to the formulas so viewers without a calculation engine show them."""
    zin = zipfile.ZipFile(path)
    wbx = zin.read("xl/workbook.xml").decode("utf-8")
    rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    target = {}
    for name, rid in re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        m = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"', rels) or re.search(rf'Target="([^"]+)"[^>]*Id="{rid}"', rels)
        target[name] = "xl/" + m.group(1).lstrip("/").replace("xl/", "")
    by_file = collections.defaultdict(dict)
    for (sheet, addr), v in cached.items():
        by_file[target[sheet]][addr] = v
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            vals = by_file.get(item.filename)
            if vals:
                def repl(m):
                    attrs, formula = m.group(1), m.group(2)
                    addr = re.search(r'r="([A-Z]+\d+)"', attrs).group(1)
                    if addr not in vals:
                        return m.group(0)
                    v = vals[addr]
                    attrs = re.sub(r'\st="[^"]*"', "", attrs)
                    if isinstance(v, str):
                        return f'<c{attrs} t="str"><f>{formula}</f><v>{v}</v></c>'
                    return f"<c{attrs}><f>{formula}</f><v>{v}</v></c>"
                data = re.sub(r"<c(\s[^>]*?)>\s*<f>(.*?)</f>\s*(?:<v\s*/>|<v>.*?</v>)?\s*</c>", repl,
                              data.decode("utf-8"), flags=re.S).encode("utf-8")
            zout.writestr(item, data)
    zin.close()
    os.replace(tmp, path)
