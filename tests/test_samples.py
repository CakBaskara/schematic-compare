"""End-to-end run on the Joystick example PDFs; skipped when they are not on this machine."""
import os

import openpyxl
import pytest

from sch_compare import cli, pdfdoc, pins

BEFORE = os.environ.get("SCH_COMPARE_BEFORE", r"C:/Users/ACER/Downloads/JOYSTICK_EVB_KC_V01_04091617.pdf")
AFTER = os.environ.get("SCH_COMPARE_AFTER", r"C:/Users/ACER/Downloads/joystick2_evb_kc_v01_0715.pdf")
CONFIG = os.path.join(os.path.dirname(__file__), "..", "examples", "joystick_vs_joystick2.json")

pytestmark = pytest.mark.skipif(not (os.path.isfile(BEFORE) and os.path.isfile(AFTER)), reason="example PDFs not found")


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    out = tmp_path_factory.mktemp("out") / "cpl.xlsx"
    return cli.run(AFTER, BEFORE, output=str(out), config=CONFIG, auto_order=True, log=lambda *a: None)


def test_order_and_part_counts(result):
    assert result["before"]["doc"].name == os.path.basename(BEFORE)
    assert len(result["before"]["parts"]) == 384
    assert len(result["after"]["parts"]) == 514


def test_known_changes(result):
    rows = {r.ref: r for r in result["rows"]}
    assert rows["UA2"].kind == "Removed" and rows["UA2"].b_val == "SS88C9H"
    assert rows["UA7"].kind == "Added" and rows["UA7"].a_val == "AW32001A"
    assert (rows["R18"].b_val, rows["R18"].a_val) == ("2.2K", "4.7R")
    assert rows["RZA113"].a_val == "NC" and rows["RZA113"].mount == "Not Mounted"
    assert rows["CZA8"].a_val == "10UF 16V X5R"
    assert rows["Q1"].kind == "Moved"
    assert rows["U17"].a_val == "PY32F040R1BT6-LQFP-64"


def test_blocks_found(result):
    after = {b.name for b in result["after"]["blocks"].values()}
    before = {b.name for b in result["before"]["blocks"].values()}
    assert {"IMU", "DCIN OVP", "Trigger/Sub/Shoulder Keys", "Back Buttons", "I2C Level Shifter (untitled)"} <= after
    assert {"ANT", "MAIN", "Headphone amp", "Mappable Action Buttons"} <= before
    imu = next(b for b in result["after"]["blocks"].values() if b.name == "IMU")
    assert "U4" in imu.refs and imu.region[2] < 440


def test_connector_pins():
    nb = pins.connector_nets(pdfdoc.load(BEFORE), "J1")
    na = pins.connector_nets(pdfdoc.load(AFTER), "J1")
    rows = {n: (o, w) for n, o, w in pins.pin_rows(nb, na)}
    assert len(rows) == 40
    assert rows[46] == ("B_SCREEN", "GPIO56")
    assert rows[88] == ("(not connected)", "DC5V")
    assert 72 not in rows and 37 not in rows


def test_workbook(result):
    wb = openpyxl.load_workbook(result["output"], data_only=True)
    assert wb.sheetnames == ["Overview", "Change Point List", "Circuit Diagram Change", "Pin Change J1", "Extraction Review"]
    cpl = wb["Change Point List"]
    assert cpl["J11"].value in ("Mounted", "Not Mounted")
    ov = wb["Overview"]
    counts = [c.value for row in ov.iter_rows(min_col=6, max_col=6) for c in row if isinstance(c.value, int)]
    assert len(result["rows"]) in counts
